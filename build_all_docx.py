#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Сборка трёх Word-документов Модуля 1: версия ученицы, версия учителя, задачник с промтами."""
import re, os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

BASE = "/home/user/Amc8/amc8/modul-1-chisla-detektivy"
OUTDIR = "/home/user/Amc8"

SUP = str.maketrans("0123456789+-=()n", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿ")
SUB = str.maketrans("0123456789+-=()", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎")
def sup(s): return s.translate(SUP) if all(c in "0123456789+-=()n" for c in s) else "^(" + s + ")"
def sub(s): return s.translate(SUB) if all(c in "0123456789+-=()" for c in s) else "_(" + s + ")"

def latex(s):
    s = re.sub(r"\\text\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\mathbf\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\textbf\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\overline\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\sqrt\{([^{}]*)\}", r"√(\1)", s)
    s = re.sub(r"\\sqrt\s*(\w)", r"√\1", s)
    s = re.sub(r"\\pmod\s*\{([^{}]*)\}", r" (mod \1)", s)
    s = re.sub(r"\\pmod\s*(\w+)", r" (mod \1)", s)
    s = re.sub(r"\\bmod", " mod ", s)
    for _ in range(3):
        s = re.sub(r"\\[dt]?frac\{([^{}]*)\}\{([^{}]*)\}", r"(\1)/(\2)", s)
    s = s.replace("\\{", "\x01").replace("\\}", "\x02")
    s = s.replace("\\\\", " ").replace("\\ ", " ")
    repl = {
        r"\\Rightarrow": "⇒", r"\\implies": "⇒", r"\\Longrightarrow": "⇒",
        r"\\rightarrow": "→", r"\\to(?![a-z])": "→", r"\\leftarrow": "←",
        r"\\Leftrightarrow": "⇔", r"\\leftrightarrow": "↔", r"\\mapsto": "↦",
        r"\\cdot": "·", r"\\times": "×", r"\\div": "÷",
        r"\\pm(?![a-z])": "±", r"\\mp(?![a-z])": "∓", r"\\geq": "≥", r"\\ge(?![a-z])": "≥",
        r"\\leq": "≤", r"\\le(?![a-z])": "≤", r"\\neq": "≠", r"\\approx": "≈",
        r"\\equiv": "≡", r"\\pi": "π", r"\\angle": "∠", r"\\parallel": "∥",
        r"\\cap": "∩", r"\\cup": "∪", r"\\in": "∈", r"\\dots": "…",
        r"\\ldots": "…", r"\\cdots": "…", r"\\lceil": "⌈", r"\\rceil": "⌉",
        r"\\lfloor": "⌊", r"\\rfloor": "⌋", r"\\min": "min", r"\\max": "max",
        r"\\left": "", r"\\right": "", r"\\quad": "  ", r"\\qquad": "    ", r"\\%": "%",
        r"\\,": " ", r"\\;": " ", r"\\!": "",
    }
    for k, v in repl.items():
        s = re.sub(k, v, s)
    s = re.sub(r"\^\{([^{}]*)\}", lambda m: sup(m.group(1)), s)
    s = re.sub(r"\^(\w)", lambda m: sup(m.group(1)), s)
    s = re.sub(r"_\{([^{}]*)\}", lambda m: sub(m.group(1)), s)
    s = re.sub(r"_(\w)", lambda m: sub(m.group(1)), s)
    s = re.sub(r"\\([a-zA-Z]+)", r"\1", s)
    s = s.replace("{", "").replace("}", "")
    return s.replace("\x01", "{").replace("\x02", "}")

TOKEN = re.compile(r"(\$\$.*?\$\$|\$.*?\$|\*\*.*?\*\*|`[^`]*`|\*[^*]+\*)", re.S)
PH = "\x00"
def _t(s): return s.replace(PH, "$")

def emit(p, text, bold=False, italic=False):
    for part in TOKEN.split(text):
        if not part: continue
        if part.startswith("$$") and part.endswith("$$"):
            r = p.add_run(latex(part[2:-2].strip())); r.italic = True; r.bold = bold
        elif part.startswith("$") and part.endswith("$") and len(part) > 1:
            r = p.add_run(latex(part[1:-1])); r.italic = True; r.bold = bold
        elif part.startswith("**") and part.endswith("**"):
            emit(p, part[2:-2], bold=True, italic=italic)
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(_t(part[1:-1])); r.font.name = "Consolas"; r.bold = bold; r.italic = italic
        elif len(part) > 1 and part.startswith("*") and part.endswith("*"):
            emit(p, part[1:-1], bold=bold, italic=True)
        else:
            r = p.add_run(_t(part)); r.bold = bold; r.italic = italic

def add_runs(p, text): emit(p, text.replace("\\$", PH))

def is_sep(cells): return all(re.match(r"^[\s:-]*$", c) for c in cells) and any("-" in c for c in cells)

def flush_table(doc, rows):
    data = []
    for r in rows:
        cells = [c.strip() for c in r.strip().strip("|").split("|")]
        if is_sep(cells): continue
        data.append(cells)
    if not data: return
    ncol = max(len(r) for r in data)
    t = doc.add_table(rows=len(data), cols=ncol); t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, r in enumerate(data):
        for j in range(ncol):
            c = t.cell(i, j); c.paragraphs[0].text = ""
            add_runs(c.paragraphs[0], (r[j] if j < len(r) else "").strip())
            if i == 0:
                for run in c.paragraphs[0].runs: run.bold = True
    doc.add_paragraph()

def render_text(doc, content):
    lines = content.split("\n")
    i, in_code, code_buf, tbl = 0, False, [], []
    while i < len(lines):
        ln = lines[i]
        if ln.strip().startswith("```"):
            if in_code:
                pre = doc.add_paragraph(); run = pre.add_run("\n".join(code_buf))
                run.font.name = "Consolas"; run.font.size = Pt(9); run.font.color.rgb = RGBColor(0x33,0x33,0x33)
                code_buf, in_code = [], False
            else: in_code = True
            i += 1; continue
        if in_code: code_buf.append(ln); i += 1; continue
        if ln.strip().startswith("|"): tbl.append(ln); i += 1; continue
        elif tbl: flush_table(doc, tbl); tbl = []
        s = ln.rstrip()
        if not s.strip(): i += 1; continue
        if re.match(r"^(---+|\*\*\*+)$", s.strip()):
            doc.add_paragraph().add_run("─"*30).font.color.rgb = RGBColor(0xAA,0xAA,0xAA); i += 1; continue
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            h = doc.add_heading(level=min(len(m.group(1)),4)); add_runs(h, m.group(2)); i += 1; continue
        if s.lstrip().startswith(">"):
            p = doc.add_paragraph(style="Intense Quote"); add_runs(p, re.sub(r"^\s*>\s?","",s)); i += 1; continue
        mm = re.match(r"^(\s*)([-*])\s+(.*)$", s)
        if mm:
            p = doc.add_paragraph(style="List Bullet"); add_runs(p, mm.group(3)); i += 1; continue
        mn = re.match(r"^(\s*)\d+\.\s+(.*)$", s)
        if mn:
            p = doc.add_paragraph(style="List Number"); add_runs(p, mn.group(2)); i += 1; continue
        p = doc.add_paragraph(); add_runs(p, s); i += 1
    if tbl: flush_table(doc, tbl)

def render_file(doc, path): render_text(doc, open(path, encoding="utf-8").read())

def new_doc():
    doc = Document(); st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    return doc

def title_page(doc, subtitle):
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("AMC 8 · Подготовка"); r.bold=True; r.font.size=Pt(16); r.font.color.rgb=RGBColor(0x2E,0x4B,0x8E)
    h = doc.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run("МОДУЛЬ 1 · Числа-детективы"); r.bold=True; r.font.size=Pt(26)
    s = doc.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run(subtitle); r.italic=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor(0x55,0x55,0x55)
    doc.add_page_break()

LESSONS = ["urok-1-delimost","urok-2-nod-nok","urok-3-ostatki-chetnost","urok-4-poslednyaya-cifra","urok-5-povtorenie"]
LNAMES = ["Урок 1 · Делимость и признаки","Урок 2 · НОД и НОК","Урок 3 · Остатки и чётность",
          "Урок 4 · Последняя цифра степени","Урок 5 · Повторение"]
ALL8 = ["01-plan-i-konspekt-uchitelya.md","02-materialy-uchenice.md","03-zadachi.md",
        "04-resheniya-i-podskazki.md","05-prezentaciya.md","06-promty-kartinok.md",
        "07-tetrad-oshibok.md","08-domashka.md"]

# ---------- clean for student ----------
def clean_student(text, is_zachet=False):
    if is_zachet:
        idx = text.find("## [ДЛЯ УЧИТЕЛЯ]")
        if idx >= 0: text = text[:idx]
    out = []
    for ln in text.split("\n"):
        if "[ДЛЯ УЧИТЕЛЯ]" in ln: continue
        if ln.strip().startswith(">") and ("Доступно:" in ln or "[ДЛЯ" in ln): continue
        out.append(ln)
    return "\n".join(out)

# ---------- 1) STUDENT ----------
def build_student():
    doc = new_doc()
    title_page(doc, "Версия для ученицы\nтеория · задачи · домашки (без ответов)")
    for slug, name in zip(LESSONS, LNAMES):
        doc.add_heading(name, level=1)
        for f in ["02-materialy-uchenice.md","03-zadachi.md","08-domashka.md"]:
            p = os.path.join(BASE, slug, f)
            if os.path.exists(p): render_text(doc, clean_student(open(p,encoding="utf-8").read()))
        doc.add_page_break()
    doc.add_heading("Зачёт модуля", level=1)
    render_text(doc, clean_student(open(os.path.join(BASE,"zachet-modulya.md"),encoding="utf-8").read(), is_zachet=True))
    out = os.path.join(OUTDIR, "Модуль-1-ученица.docx"); doc.save(out); return out, doc

# ---------- 2) TEACHER ----------
def build_teacher():
    doc = new_doc()
    title_page(doc, "Версия для учителя — полный комплект\nконспекты · решения · презентации · промты · тетрадь ошибок · зачёт")
    doc.add_heading("Обзор модуля", level=1); render_file(doc, os.path.join(BASE,"README.md")); doc.add_page_break()
    for slug, name in zip(LESSONS, LNAMES):
        doc.add_heading(name, level=1)
        for f in ALL8:
            p = os.path.join(BASE, slug, f)
            if os.path.exists(p): render_file(doc, p)
        doc.add_page_break()
    doc.add_heading("Зачёт модуля", level=1); render_file(doc, os.path.join(BASE,"zachet-modulya.md"))
    out = os.path.join(OUTDIR, "Модуль-1-учитель.docx"); doc.save(out); return out, doc

# ---------- parse helpers for задачник ----------
def slides_of(path):
    txt = open(path, encoding="utf-8").read()
    out = []
    for blk in re.split(r"^### ", txt, flags=re.M)[1:]:
        head = re.search(r"\*\*Заголовок:\*\*\s*(.+)", blk)
        scr = re.search(r"\*\*На слайде:\*\*\s*(.+)", blk)
        if head and scr: out.append((head.group(1).strip(), scr.group(1).strip()))
    return out

def problems_lesson(path):
    txt = open(path, encoding="utf-8").read()
    parts = re.split(r"^## (Задача[^\n]*)$", txt, flags=re.M)
    res = []
    for i in range(1, len(parts), 2):
        head = parts[i].strip(); body = parts[i+1].strip()
        res.append((head, body))
    return res

def problems_zachet(path):
    txt = open(path, encoding="utf-8").read()
    seg = txt.split("## [ДЛЯ УЧЕНИЦЫ]")[1].split("## [ДЛЯ УЧИТЕЛЯ]")[0]
    res = []
    for m in re.finditer(r"\*\*(З\d+)[^\n]*\*\*\s*([^\n]*)\n(\$[^\n]*\$)", seg):
        res.append((m.group(1), m.group(2).strip() + "\n\n" + m.group(3).strip()))
    # fallback simpler split
    if not res:
        for blk in re.split(r"(?=^\*\*З\d+)", seg, flags=re.M):
            blk = blk.strip()
            if blk.startswith("**З"):
                res.append(("", blk))
    return res

def parse_prompts(path):
    txt = open(path, encoding="utf-8").read()
    data = {}
    for sm in re.finditer(r"^## (U\d+|ZACHET)\s*(.*?)(?=^## |\Z)", txt, flags=re.M|re.S):
        key = sm.group(1); body = sm.group(2)
        items = {}
        for zm in re.finditer(r"^### Z(\d+)\s*(.*?)(?=^### |\Z)", body, flags=re.M|re.S):
            items[int(zm.group(1))] = zm.group(2).strip()
        data[key] = items
    return data

# ---------- 3) ЗАДАЧНИК ----------
def build_zadachnik():
    doc = new_doc()
    title_page(doc, "Задачник с картинками\nтексты презентаций · задачи · промты картинок (кошки · додзё · Ever After High)")
    prompts = parse_prompts(os.path.join(BASE, "kreativnye-promty-kartinok.md"))
    for li, (slug, name) in enumerate(zip(LESSONS, LNAMES), start=1):
        doc.add_heading(name, level=1)
        # narrative from slides
        doc.add_heading("Сюжет урока (по слайдам презентации)", level=2)
        for hd, sc in slides_of(os.path.join(BASE, slug, "05-prezentaciya.md")):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, "**" + hd + ".** " + sc)
        # problems + prompts
        doc.add_heading("Задачи и промты картинок", level=2)
        probs = problems_lesson(os.path.join(BASE, slug, "03-zadachi.md"))
        pkey = prompts.get("U%d" % li, {})
        for idx, (head, body) in enumerate(probs, start=1):
            h = doc.add_heading(level=3); add_runs(h, head)
            render_text(doc, body)
            if idx in pkey:
                pp = doc.add_paragraph(); r = pp.add_run("🎨 Промт для картинки")
                r.bold = True; r.font.color.rgb = RGBColor(0xB5, 0x4A, 0x00)
                render_text(doc, pkey[idx])
            doc.add_paragraph()
        doc.add_page_break()
    # zachet
    doc.add_heading("Зачёт модуля · задачи и промты", level=1)
    zprob = problems_zachet(os.path.join(BASE, "zachet-modulya.md"))
    zkey = prompts.get("ZACHET", {})
    for idx, (head, body) in enumerate(zprob, start=1):
        h = doc.add_heading(level=3); add_runs(h, "Зачёт · Задача %d" % idx)
        render_text(doc, body)
        if idx in zkey:
            pp = doc.add_paragraph(); r = pp.add_run("🎨 Промт для картинки")
            r.bold = True; r.font.color.rgb = RGBColor(0xB5, 0x4A, 0x00)
            render_text(doc, zkey[idx])
        doc.add_paragraph()
    out = os.path.join(OUTDIR, "Модуль-1-Задачник.docx"); doc.save(out); return out, doc

for fn in (build_student, build_teacher, build_zadachnik):
    out, doc = fn()
    print("OK", os.path.basename(out), "| абзацев:", len(doc.paragraphs), "| таблиц:", len(doc.tables))
