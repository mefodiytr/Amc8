#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Универсальный сборщик Word-документов модуля: ученица / учитель / задачник.
Использование: python3 build_module.py <modul-slug>
Авто-разбор уроков из README модуля. Задачник строится, если есть kreativnye-promty-kartinok.md."""
import re, os, sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = "/home/user/Amc8/amc8"
OUTDIR = "/home/user/Amc8"

SUP = str.maketrans("0123456789+-=()n", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿ")
SUB = str.maketrans("0123456789+-=()", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎")
def sup(s): return s.translate(SUP) if all(c in "0123456789+-=()n" for c in s) else "^(" + s + ")"
def sub(s): return s.translate(SUB) if all(c in "0123456789+-=()" for c in s) else "_(" + s + ")"

def latex(s):
    s = re.sub(r"\\text\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\mathbf\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\textbf\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\overline\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\sqrt\{([^{}]*)\}", r"√(\1)", s)
    s = re.sub(r"\\sqrt\s*(\w)", r"√\1", s)
    s = re.sub(r"\\pmod\s*\{([^{}]*)\}", r" (mod \1)", s)
    s = re.sub(r"\\pmod\s*(\w+)", r" (mod \1)", s)
    s = re.sub(r"\\bmod", " mod ", s)
    for _ in range(3):
        s = re.sub(r"\\[dt]?frac\{([^{}]*)\}\{([^{}]*)\}", r"(\1)/(\2)", s)
    s = s.replace("\\{", "\x01").replace("\\}", "\x02")
    s = s.replace("\\\\", " ").replace("\\ ", " ")
    repl = {
        r"\\Rightarrow": "⇒", r"\\implies": "⇒", r"\\Longrightarrow": "⇒",
        r"\\rightarrow": "→", r"\\to(?![a-z])": "→", r"\\leftarrow": "←",
        r"\\Leftrightarrow": "⇔", r"\\leftrightarrow": "↔", r"\\mapsto": "↦",
        r"\\cdot": "·", r"\\times": "×", r"\\div": "÷",
        r"\\pm(?![a-z])": "±", r"\\mp(?![a-z])": "∓", r"\\geq": "≥", r"\\ge(?![a-z])": "≥",
        r"\\leq": "≤", r"\\le(?![a-z])": "≤", r"\\neq": "≠", r"\\approx": "≈",
        r"\\equiv": "≡", r"\\pi": "π", r"\\angle": "∠", r"\\parallel": "∥",
        r"\\cap": "∩", r"\\cup": "∪", r"\\in": "∈", r"\\dots": "…",
        r"\\ldots": "…", r"\\cdots": "…", r"\\lceil": "⌈", r"\\rceil": "⌉",
        r"\\lfloor": "⌊", r"\\rfloor": "⌋", r"\\min": "min", r"\\max": "max",
        r"\\left": "", r"\\right": "", r"\\quad": "  ", r"\\qquad": "    ", r"\\%": "%",
        r"\\,": " ", r"\\;": " ", r"\\!": "",
    }
    for k, v in repl.items():
        s = re.sub(k, v, s)
    s = re.sub(r"\^\{([^{}]*)\}", lambda m: sup(m.group(1)), s)
    s = re.sub(r"\^(\w)", lambda m: sup(m.group(1)), s)
    s = re.sub(r"_\{([^{}]*)\}", lambda m: sub(m.group(1)), s)
    s = re.sub(r"_(\w)", lambda m: sub(m.group(1)), s)
    s = re.sub(r"\\([a-zA-Z]+)", r"\1", s)
    s = s.replace("{", "").replace("}", "")
    return s.replace("\x01", "{").replace("\x02", "}")

TOKEN = re.compile(r"(\$\$.*?\$\$|\$.*?\$|\*\*.*?\*\*|`[^`]*`|\*[^*]+\*)", re.S)
PH = "\x00"
def _t(s): return s.replace(PH, "$")
def emit(p, text, bold=False, italic=False):
    for part in TOKEN.split(text):
        if not part: continue
        if part.startswith("$$") and part.endswith("$$"):
            r = p.add_run(latex(part[2:-2].strip())); r.italic=True; r.bold=bold
        elif part.startswith("$") and part.endswith("$") and len(part) > 1:
            r = p.add_run(latex(part[1:-1])); r.italic=True; r.bold=bold
        elif part.startswith("**") and part.endswith("**"):
            emit(p, part[2:-2], bold=True, italic=italic)
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(_t(part[1:-1])); r.font.name="Consolas"; r.bold=bold; r.italic=italic
        elif len(part) > 1 and part.startswith("*") and part.endswith("*"):
            emit(p, part[1:-1], bold=bold, italic=True)
        else:
            r = p.add_run(_t(part)); r.bold=bold; r.italic=italic
def add_runs(p, text): emit(p, text.replace("\\$", PH))
def is_sep(cells): return all(re.match(r"^[\s:-]*$", c) for c in cells) and any("-" in c for c in cells)
def flush_table(doc, rows):
    data=[]
    for r in rows:
        cells=[c.strip() for c in r.strip().strip("|").split("|")]
        if is_sep(cells): continue
        data.append(cells)
    if not data: return
    ncol=max(len(r) for r in data)
    t=doc.add_table(rows=len(data),cols=ncol); t.style="Light Grid Accent 1"; t.alignment=WD_TABLE_ALIGNMENT.LEFT
    for i,r in enumerate(data):
        for j in range(ncol):
            c=t.cell(i,j); c.paragraphs[0].text=""
            add_runs(c.paragraphs[0],(r[j] if j<len(r) else "").strip())
            if i==0:
                for run in c.paragraphs[0].runs: run.bold=True
    doc.add_paragraph()
def render_text(doc, content):
    lines=content.split("\n"); i=0; in_code=False; code_buf=[]; tbl=[]
    while i<len(lines):
        ln=lines[i]
        if ln.strip().startswith("```"):
            if in_code:
                pre=doc.add_paragraph(); run=pre.add_run("\n".join(code_buf))
                run.font.name="Consolas"; run.font.size=Pt(9); run.font.color.rgb=RGBColor(0x33,0x33,0x33)
                code_buf=[]; in_code=False
            else: in_code=True
            i+=1; continue
        if in_code: code_buf.append(ln); i+=1; continue
        if ln.strip().startswith("|"): tbl.append(ln); i+=1; continue
        elif tbl: flush_table(doc,tbl); tbl=[]
        s=ln.rstrip()
        if not s.strip(): i+=1; continue
        if re.match(r"^(---+|\*\*\*+)$", s.strip()):
            doc.add_paragraph().add_run("─"*30).font.color.rgb=RGBColor(0xAA,0xAA,0xAA); i+=1; continue
        m=re.match(r"^(#{1,6})\s+(.*)$", s)
        if m: h=doc.add_heading(level=min(len(m.group(1)),4)); add_runs(h,m.group(2)); i+=1; continue
        if s.lstrip().startswith(">"):
            p=doc.add_paragraph(style="Intense Quote"); add_runs(p,re.sub(r"^\s*>\s?","",s)); i+=1; continue
        mm=re.match(r"^(\s*)([-*])\s+(.*)$", s)
        if mm: p=doc.add_paragraph(style="List Bullet"); add_runs(p,mm.group(3)); i+=1; continue
        mn=re.match(r"^(\s*)\d+\.\s+(.*)$", s)
        if mn: p=doc.add_paragraph(style="List Number"); add_runs(p,mn.group(2)); i+=1; continue
        p=doc.add_paragraph(); add_runs(p,s); i+=1
    if tbl: flush_table(doc,tbl)
def render_file(doc, path): render_text(doc, open(path,encoding="utf-8").read())
def new_doc():
    doc=Document(); st=doc.styles["Normal"]; st.font.name="Calibri"; st.font.size=Pt(11); return doc
def title_page(doc, modtitle, subtitle):
    t=doc.add_paragraph(); t.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=t.add_run("AMC 8 · Подготовка"); r.bold=True; r.font.size=Pt(16); r.font.color.rgb=RGBColor(0x2E,0x4B,0x8E)
    h=doc.add_paragraph(); h.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=h.add_run(modtitle); r.bold=True; r.font.size=Pt(24)
    s=doc.add_paragraph(); s.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=s.add_run(subtitle); r.italic=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor(0x55,0x55,0x55)
    doc.add_page_break()

def parse_readme(mod_dir):
    txt=open(os.path.join(mod_dir,"README.md"),encoding="utf-8").read()
    title=re.search(r"^#\s+(.*)$", txt, flags=re.M).group(1).strip()
    lessons=[]
    for m in re.finditer(r"^\|\s*(\d+)\s*\|\s*([^|]+?)\s*\|\s*`([^`]+)`", txt, flags=re.M):
        lessons.append((int(m.group(1)), m.group(2).strip(), m.group(3).strip()))
    return title, lessons

ALL8=["01-plan-i-konspekt-uchitelya.md","02-materialy-uchenice.md","03-zadachi.md",
      "04-resheniya-i-podskazki.md","05-prezentaciya.md","06-promty-kartinok.md",
      "07-tetrad-oshibok.md","08-domashka.md"]

def clean_student(text, is_zachet=False):
    if is_zachet:
        idx=text.find("## [ДЛЯ УЧИТЕЛЯ]")
        if idx>=0: text=text[:idx]
    out=[]
    for ln in text.split("\n"):
        if "[ДЛЯ УЧИТЕЛЯ]" in ln: continue
        if ln.strip().startswith(">") and ("Доступно:" in ln or "[ДЛЯ" in ln): continue
        out.append(ln)
    return "\n".join(out)

def slides_of(path):
    if not os.path.exists(path): return []
    txt=open(path,encoding="utf-8").read(); out=[]
    for blk in re.split(r"^### ", txt, flags=re.M)[1:]:
        head=re.search(r"\*\*Заголовок:\*\*\s*(.+)", blk)
        scr=re.search(r"\*\*На слайде:\*\*\s*(.+)", blk)
        if head and scr: out.append((head.group(1).strip(), scr.group(1).strip()))
    return out

def problems_lesson(path):
    txt=open(path,encoding="utf-8").read()
    parts=re.split(r"^## (Задача[^\n]*)$", txt, flags=re.M)
    return [(parts[i].strip(), parts[i+1].strip()) for i in range(1,len(parts),2)]

def problems_zachet(path):
    txt=open(path,encoding="utf-8").read()
    seg=txt.split("## [ДЛЯ УЧЕНИЦЫ]")[1].split("## [ДЛЯ УЧИТЕЛЯ]")[0]
    res=[]
    for blk in re.split(r"(?=^\*\*З\d+)", seg, flags=re.M):
        blk=blk.strip()
        if blk.startswith("**З"): res.append(blk)
    return res

def parse_prompts(path):
    if not os.path.exists(path): return {}
    txt=open(path,encoding="utf-8").read(); data={}
    for sm in re.finditer(r"^## (U\d+|ZACHET)\s*(.*?)(?=^## |\Z)", txt, flags=re.M|re.S):
        items={}
        for zm in re.finditer(r"^### Z(\d+)\s*(.*?)(?=^### |\Z)", sm.group(2), flags=re.M|re.S):
            items[int(zm.group(1))]=zm.group(2).strip()
        data[sm.group(1)]=items
    return data

def build(mod_slug):
    mod_dir=os.path.join(ROOT, mod_slug)
    title, lessons=parse_readme(mod_dir)
    short=title.split("·")[-1].strip() if "·" in title else title
    modnum=re.search(r"Модул[ья]\s*(\d+)", title)
    tag="Модуль-%s" % (modnum.group(1) if modnum else mod_slug)
    made=[]

    # 1) УЧЕНИЦА
    doc=new_doc(); title_page(doc, title, "Версия для ученицы — теория · задачи · домашки (без ответов)")
    for num,name,slug in lessons:
        doc.add_heading("Урок %d · %s" % (num,name), level=1)
        for f in ["02-materialy-uchenice.md","03-zadachi.md","08-domashka.md"]:
            p=os.path.join(mod_dir,slug,f)
            if os.path.exists(p): render_text(doc, clean_student(open(p,encoding="utf-8").read()))
        doc.add_page_break()
    zpath=os.path.join(mod_dir,"zachet-modulya.md")
    if os.path.exists(zpath):
        doc.add_heading("Зачёт модуля", level=1)
        render_text(doc, clean_student(open(zpath,encoding="utf-8").read(), True))
    o=os.path.join(OUTDIR, tag+"-ученица.docx"); doc.save(o); made.append((o,doc))

    # 2) УЧИТЕЛЬ
    doc=new_doc(); title_page(doc, title, "Версия для учителя — полный комплект")
    doc.add_heading("Обзор модуля", level=1); render_file(doc, os.path.join(mod_dir,"README.md")); doc.add_page_break()
    for num,name,slug in lessons:
        doc.add_heading("Урок %d · %s" % (num,name), level=1)
        for f in ALL8:
            p=os.path.join(mod_dir,slug,f)
            if os.path.exists(p): render_file(doc,p)
        doc.add_page_break()
    if os.path.exists(zpath):
        doc.add_heading("Зачёт модуля", level=1); render_file(doc, zpath)
    o=os.path.join(OUTDIR, tag+"-учитель.docx"); doc.save(o); made.append((o,doc))

    # 3) ЗАДАЧНИК (если есть промты)
    prompts=parse_prompts(os.path.join(mod_dir,"kreativnye-promty-kartinok.md"))
    if prompts:
        doc=new_doc(); title_page(doc, title, "Задачник с картинками — тексты презентаций · задачи · промты (кошки · додзё · Ever After High)")
        for li,(num,name,slug) in enumerate(lessons, start=1):
            doc.add_heading("Урок %d · %s" % (num,name), level=1)
            sl=slides_of(os.path.join(mod_dir,slug,"05-prezentaciya.md"))
            if sl:
                doc.add_heading("Сюжет урока (по слайдам презентации)", level=2)
                for hd,sc in sl:
                    add_runs(doc.add_paragraph(style="List Bullet"), "**"+hd+".** "+sc)
            doc.add_heading("Задачи и промты картинок", level=2)
            probs=problems_lesson(os.path.join(mod_dir,slug,"03-zadachi.md"))
            pk=prompts.get("U%d"%num, {})
            for idx,(head,body) in enumerate(probs, start=1):
                h=doc.add_heading(level=3); add_runs(h,head); render_text(doc,body)
                if idx in pk:
                    r=doc.add_paragraph().add_run("🎨 Промт для картинки"); r.bold=True; r.font.color.rgb=RGBColor(0xB5,0x4A,0x00)
                    render_text(doc, pk[idx])
                doc.add_paragraph()
            doc.add_page_break()
        if os.path.exists(zpath):
            doc.add_heading("Зачёт модуля · задачи и промты", level=1)
            zk=prompts.get("ZACHET", {})
            for idx,body in enumerate(problems_zachet(zpath), start=1):
                h=doc.add_heading(level=3); add_runs(h,"Зачёт · Задача %d"%idx); render_text(doc,body)
                if idx in zk:
                    r=doc.add_paragraph().add_run("🎨 Промт для картинки"); r.bold=True; r.font.color.rgb=RGBColor(0xB5,0x4A,0x00)
                    render_text(doc, zk[idx])
                doc.add_paragraph()
        o=os.path.join(OUTDIR, tag+"-Задачник.docx"); doc.save(o); made.append((o,doc))

    for o,d in made:
        print("OK", os.path.basename(o), "| абзацев:", len(d.paragraphs), "| таблиц:", len(d.tables))

if __name__ == "__main__":
    build(sys.argv[1])
