#!/usr/bin/env python3
"""
validate_deck.py — программная проверка собранной презентации урока.
Не заменяет просмотр глазами, но ловит частые баги.

    python tools/validate_deck.py <путь>/deck.pptx

Проверяет:
  - текст не вылезает за поля слайда (мин. отступ >= 2%);
  - картинки вшиты (media);
  - у каждого слайда есть заметки докладчику;
  - предупреждение «окно, не экран» встречается в заметках.
Требует: pip install python-pptx
"""
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Emu
except ImportError:
    sys.exit("Нужен python-pptx:  pip install python-pptx")


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    p = Path(sys.argv[1])
    if not p.exists():
        sys.exit(f"нет файла: {p}")

    pres = Presentation(str(p))
    W, H = pres.slide_width, pres.slide_height
    n = len(pres.slides)
    print(f"слайдов: {n}")

    problems = 0
    no_notes = []
    off_slide = []
    window_warning = False

    for i, slide in enumerate(pres.slides, 1):
        # заметки
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text or ""
        if not notes.strip():
            no_notes.append(i)
        low = notes.lower()
        if ("окно" in low and "экран" in low) or ("window" in low):
            window_warning = True

        # текст за полями
        for sh in slide.shapes:
            if not sh.has_text_frame:
                continue
            if not (sh.text_frame.text or "").strip():
                continue
            try:
                l, t = sh.left, sh.top
                r, b = l + sh.width, t + sh.height
            except TypeError:
                continue
            mL, mR = l / W, (W - r) / W
            mT, mB = t / H, (H - b) / H
            if min(mL, mR, mT, mB) < 0.02:
                off_slide.append((i, round(min(mL, mR, mT, mB) * 100, 1)))

    media = [part for part in pres.part.package.iter_parts()
             if "media" in part.partname]
    print(f"картинок вшито: {len(media)}")

    if no_notes:
        print(f"⚠ без заметок докладчику: слайды {no_notes}")
        problems += 1
    else:
        print("✓ заметки есть на всех слайдах")

    if off_slide:
        print(f"⚠ текст у края (<2%): {off_slide}")
        problems += 1
    else:
        print("✓ текст в полях")

    if window_warning:
        print("✓ предупреждение «окно, не экран» найдено в заметках")
    else:
        print("⚠ не найдено предупреждение «шарить окно, не экран» — добавь в заметки титула")
        problems += 1

    print()
    print("НАПОМИНАНИЕ: глазами проверь то, что скрипт не видит — "
          "перепутанные слои/цвета фигур, числа на диаграммах = числам в условии, "
          "панель не наехала на героиню.")
    print("ИТОГ:", "ок" if problems == 0 else f"замечаний: {problems}")
    sys.exit(0 if problems == 0 else 1)


if __name__ == "__main__":
    main()
