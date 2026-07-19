#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Конвертер одного markdown-файла в читаемый .docx. LaTeX ($...$) → Unicode.
Использование: python3 md2docx.py вход.md выход.docx "Заголовок" "Подзаголовок"
"""
import re, os, sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

SUP = str.maketrans("0123456789+-=()n", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁼⁽⁾ⁿ")
SUB = str.maketrans("0123456789+-=()", "₀₁₂₃₄₅₆₇₈₉₊₋₌₍₎")
def sup(s): return s.translate(SUP) if all(c in "0123456789+-=()n" for c in s) else "^(" + s + ")"
def sub(s): return s.translate(SUB) if all(c in "0123456789+-=()" for c in s) else "_(" + s + ")"

def latex(s):
    s = re.sub(r"\\text\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\mathbf\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\textbf\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\overline\{([^{}]*)\}", r"\1", s)
    s = re.sub(r"\\binom\{([^{}]*)\}\{([^{}]*)\}", r"C(\1,\2)", s)
    s = re.sub(r"\\sqrt\{([^{}]*)\}", r"√(\1)", s)
    s = re.sub(r"\\pmod\s*\{([^{}]*)\}", r" (mod \1)", s)
    for _ in range(3):
        s = re.sub(r"\\[dt]?frac\{([^{}]*)\}\{([^{}]*)\}", r"(\1)/(\2)", s)
    s = s.replace("\\{", "\x01").replace("\\}", "\x02")
    s = s.replace("\\\\", " ").replace("\\ ", " ")
    repl = {
        r"\\Rightarrow": "⇒", r"\\rightarrow": "→", r"\\to(?![a-z])": "→",
        r"\\cdot": "·", r"\\times": "×", r"\\div": "÷", r"\\approx": "≈",
        r"\\pm(?![a-z])": "±", r"\\geq": "≥", r"\\ge(?![a-z])": "≥",
        r"\\leq": "≤", r"\\le(?![a-z])": "≤", r"\\neq": "≠", r"\\equiv": "≡",
        r"\\dots": "…", r"\\ldots": "…", r"\\cdots": "…", r"\\pi": "π",
        r"\\left": "", r"\\right": "", r"\\quad": "  ", r"\\qquad": "    ",
        r"\\%": "%", r"\\,": " ", r"\\;": " ", r"\\!": "", r"\\angle": "∠",
    }
    for k, v in repl.items():
        s = re.sub(k, v, s)
    s = re.sub(r"\^\{([^{}]*)\}", lambda m: sup(m.group(1)), s)
    s = re.sub(r"\^(\w)", lambda m: sup(m.group(1)), s)
    s = re.sub(r"_\{([^{}]*)\}", lambda m: sub(m.group(1)), s)
    s = re.sub(r"_(\w)", lambda m: sub(m.group(1)), s)
    s = re.sub(r"\\([a-zA-Z]+)", r"\1", s)
    s = s.replace("{", "").replace("}", "")
    s = s.replace("\x01", "{").replace("\x02", "}")
    return s

TOKEN = re.compile(r"(\$\$.*?\$\$|\$.*?\$|\*\*.*?\*\*|`[^`]*`|\*[^*]+\*)", re.S)

def emit(p, text, bold=False, italic=False):
    for part in TOKEN.split(text):
        if not part: continue
        if part.startswith("$$") and part.endswith("$$"):
            r = p.add_run(latex(part[2:-2].strip())); r.italic = True; r.bold = bold
        elif part.startswith("$") and part.endswith("$") and len(part) > 1:
            r = p.add_run(latex(part[1:-1])); r.italic = True; r.bold = bold
        elif part.startswith("**") and part.endswith("**"):
            emit(p, part[2:-2], bold=True, italic=italic)
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1]); r.font.name = "Consolas"; r.bold = bold; r.italic = italic
        elif len(part) > 1 and part.startswith("*") and part.endswith("*"):
            emit(p, part[1:-1], bold=bold, italic=True)
        else:
            r = p.add_run(part); r.bold = bold; r.italic = italic

def add_runs(p, text): emit(p, text)

def is_sep(cells): return all(re.match(r"^[\s:-]*$", c) for c in cells) and any("-" in c for c in cells)

def flush_table(doc, rows):
    data = []
    for r in rows:
        cells = [c.strip() for c in r.strip().strip("|").split("|")]
        if is_sep(cells): continue
        data.append(cells)
    if not data: return
    ncol = max(len(r) for r in data)
    t = doc.add_table(rows=len(data), cols=ncol); t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, r in enumerate(data):
        for j in range(ncol):
            cell = t.cell(i, j); cell.paragraphs[0].text = ""
            add_runs(cell.paragraphs[0], (r[j] if j < len(r) else "").strip())
            if i == 0:
                for run in cell.paragraphs[0].runs: run.bold = True
    doc.add_paragraph()

def render_md(doc, text):
    lines = text.split("\n"); i = 0; in_code = False; code_buf = []; tbl = []
    while i < len(lines):
        ln = lines[i]
        if ln.strip().startswith("```"):
            if in_code:
                pre = doc.add_paragraph(); run = pre.add_run("\n".join(code_buf))
                run.font.name = "Consolas"; run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33); code_buf = []; in_code = False
            else: in_code = True
            i += 1; continue
        if in_code: code_buf.append(ln); i += 1; continue
        if ln.strip().startswith("|"): tbl.append(ln); i += 1; continue
        elif tbl: flush_table(doc, tbl); tbl = []
        s = ln.rstrip()
        if not s.strip(): i += 1; continue
        if re.match(r"^(---+|\*\*\*+)$", s.strip()):
            doc.add_paragraph().add_run("─" * 40).font.color.rgb = RGBColor(0xAA, 0xAA, 0xAA); i += 1; continue
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            lvl = min(len(m.group(1)), 4); h = doc.add_heading(level=lvl); add_runs(h, m.group(2)); i += 1; continue
        if s.lstrip().startswith(">"):
            p = doc.add_paragraph(style="Intense Quote"); add_runs(p, re.sub(r"^\s*>\s?", "", s)); i += 1; continue
        mm = re.match(r"^(\s*)([-*])\s+(.*)$", s)
        if mm:
            p = doc.add_paragraph(style="List Bullet"); add_runs(p, mm.group(3)); i += 1; continue
        mn = re.match(r"^(\s*)\d+\.\s+(.*)$", s)
        if mn:
            p = doc.add_paragraph(style="List Number"); add_runs(p, mn.group(2)); i += 1; continue
        p = doc.add_paragraph(); add_runs(p, s); i += 1
    if tbl: flush_table(doc, tbl)

def build(md_path, out_path, title, subtitle):
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("AMC 8 · Подготовка"); r.bold = True; r.font.size = Pt(15); r.font.color.rgb = RGBColor(0x2E, 0x4B, 0x8E)
    h = doc.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(title); r.bold = True; r.font.size = Pt(26)
    if subtitle:
        sp = doc.add_paragraph(); sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = sp.add_run(subtitle); r.italic = True; r.font.size = Pt(12); r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    doc.add_page_break()
    render_md(doc, open(md_path, encoding="utf-8").read())
    doc.save(out_path)
    print("OK:", out_path, "| параграфов:", len(doc.paragraphs), "| таблиц:", len(doc.tables))

if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2], sys.argv[3], sys.argv[4] if len(sys.argv) > 4 else "")
